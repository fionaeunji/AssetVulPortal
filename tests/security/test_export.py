"""Phase 8: Excel Export + Formula Injection 방어."""
from __future__ import annotations

import io
from datetime import datetime, timezone

import pytest
from openpyxl import Workbook, load_workbook

from app.services.exporter import _put, export_filename, sanitize_text


@pytest.mark.parametrize("payload", [
    "=HYPERLINK(\"http://evil\",\"click\")", "=cmd|' /C calc'!A0", "+1+1", "-2+3", "@SUM(1,1)",
    "\t=1+1", "\r=1+1", "\n=1", "＝1+1", "＋1", "－1", "＠1",
])
def test_formula_prefixes_neutralized(payload):
    out = sanitize_text(payload)
    assert out.startswith("'") and out[1:] == payload


@pytest.mark.parametrize("safe", ["web-dmz-01", "Apache HTTP Server", "10.0.20348.2582", "가상담당자01", " =1"])
def test_normal_values_unchanged(safe):
    assert sanitize_text(safe) == safe


def _roundtrip(values):
    wb = Workbook()
    ws = wb.active
    for i, v in enumerate(values, start=1):
        _put(ws, i, 1, v)
    buf = io.BytesIO()
    wb.save(buf)
    return load_workbook(io.BytesIO(buf.getvalue())).active


def test_saved_cells_are_text_not_formulas():
    payloads = ["=1+1", "+SUM(A1)", "-1+1", "@x", "=HYPERLINK(\"http://evil\")"]
    ws = _roundtrip(payloads)
    for i, p in enumerate(payloads, start=1):
        c = ws.cell(row=i, column=1)
        assert c.data_type == "s" and c.value == "'" + p


def test_numbers_and_dates_keep_types():
    ws = _roundtrip([9.8, 0.4, datetime(2026, 9, 28, 0, 0, tzinfo=timezone.utc), None, True])
    assert ws["A1"].value == 9.8 and ws["A2"].value == 0.4
    assert ws["A3"].value == datetime(2026, 9, 28, 9, 0)          # KST 변환
    assert ws["A4"].value is None and ws["A5"].value == "Y"


def test_long_text_truncated_to_excel_limit():
    assert len(sanitize_text("a" * 40000)) == 32000


def test_filename_is_server_generated():
    assert export_filename(datetime(2026, 9, 28, 0, 5, tzinfo=timezone.utc)) == "vuln_report_20260928_0905.xlsx"
