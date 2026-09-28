"""Phase 8: 웹 Excel Export (필터 반영, 감사로그, 외부/사용자 데이터 무해화)."""
from __future__ import annotations

import io

from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import select

from app.models import Asset, AuditLog, Vulnerability
from tests.web.test_web import login, portal  # noqa: F401  (fixture 재사용)


def _sheet(resp):
    wb = load_workbook(io.BytesIO(resp.content))
    ws = wb["취약점 현황"]
    header = [c.value for c in ws[1]]
    rows = [dict(zip(header, [c.value for c in r])) for r in ws.iter_rows(min_row=2)]
    return wb, ws, header, rows


def test_export_requires_login(portal):  # noqa: F811
    app, _, _ = portal
    r = TestClient(app).get("/export.xlsx", follow_redirects=False)
    assert r.status_code == 303 and r.headers["location"] == "/login"


def test_export_columns_filter_and_audit(portal):  # noqa: F811
    app, factory, _ = portal
    c = login(app, "viewer1")
    r = c.get("/export.xlsx?sev=긴급")
    assert r.status_code == 200
    assert r.headers["content-type"].startswith("application/vnd.openxmlformats-officedocument.spreadsheetml")
    assert 'filename="vuln_report_' in r.headers["content-disposition"]
    wb, ws, header, rows = _sheet(r)
    for col in ("등급", "CVE", "자산명", "IP", "제품", "Version", "CVSS", "Initial EPSS", "Current EPSS", "KEV",
                "자산구분", "중요도", "담당자", "탐지일", "조치기한", "남은시간(시간)", "상태", "매핑근거"):
        assert col in header
    assert rows and all(r_["등급"] == "긴급" for r_ in rows)
    web = next(r_ for r_ in rows if r_["CVE"] == "CVE-2021-41773")
    assert web["CVSS"] == 9.8 and web["Initial EPSS"] == 0.94 and web["KEV"] == "Y"
    assert "요약" in wb.sheetnames
    with factory() as s:
        log = s.execute(select(AuditLog).where(AuditLog.action == "EXPORT")).scalar_one()
        assert log.actor == "viewer1" and log.after["rows"] == len(rows) and "sev=긴급" in log.after["filter"]


def test_export_neutralizes_formula_injection_from_excel_and_nvd(portal):  # noqa: F811
    app, factory, _ = portal
    with factory() as s:
        a = s.execute(select(Asset).where(Asset.asset_code == "WEB-001")).scalar_one()
        a.name = '=HYPERLINK("http://evil.example","업데이트")'
        a.department = "+cmd|' /C calc'!A0"
        v = s.get(Vulnerability, "CVE-2021-41773")
        v.description = "@SUM(1+1)*cmd|' /C calc'!A0"
        s.commit()
    r = login(app, "viewer1").get("/export.xlsx?q=CVE-2021-41773")
    _, ws, header, rows = _sheet(r)
    row = rows[0]
    assert row["자산명"] == "'" + '=HYPERLINK("http://evil.example","업데이트")'
    assert row["부서"].startswith("'+") and row["CVE 설명"].startswith("'@")
    for cells in ws.iter_rows(min_row=2):
        for cell in cells:
            assert cell.data_type != "f"          # 수식 셀이 하나도 없어야 함
