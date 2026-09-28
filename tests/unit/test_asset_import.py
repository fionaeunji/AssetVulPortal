"""자산관리대장 Import 테스트."""
from __future__ import annotations

import io
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook
from sqlalchemy import func, select
from sqlalchemy.orm import sessionmaker

from app.models import Asset, AssetProduct, AssetVulnerability, AuditLog, ImportBatch, Vulnerability
from app.models.enums import AssessmentState, MatchType, Zone
from app.services.asset_importer import import_asset_upload, normalize_token
from app.services.vulnerability_collector import build_targets

SAMPLE = Path(__file__).resolve().parents[2] / "sample_data" / "sample_assets.xlsx"
HEAD = ["Asset ID", "자산명", "자산구분", "IP", "Vendor", "Product", "Version", "CPE", "중요도", "담당자", "부서"]
ROW = ["SRV-01", "srv01", "내부", "192.0.2.1", "Apache", "Tomcat", "9.0.30",
       "cpe:2.3:a:apache:tomcat:9.0.30:*:*:*:*:*:*:*", "상", "가상담당", "가상팀"]


@pytest.fixture()
def factory(engine):
    return sessionmaker(bind=engine, expire_on_commit=False)


def xlsx(rows, header=HEAD, formula_cells=()) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.append(header)
    for r in rows:
        ws.append(r)
    for coord, f in formula_cells:
        ws[coord] = f
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def run(factory, tmp_path, data, name="assets.xlsx"):
    return import_asset_upload(factory, filename=name, data=data, actor="op1", actor_role="operator",
                               upload_dir=tmp_path / "uploads", max_bytes=5 * 1024 * 1024)


def count(factory, model):
    with factory() as s:
        return s.execute(select(func.count()).select_from(model)).scalar()


def test_sample_file_has_20_plus_assets():
    ws = load_workbook(SAMPLE, read_only=True).active
    codes = {r[0] for r in ws.iter_rows(min_row=2, values_only=True) if r[0]}
    assert len(codes) >= 20


def test_import_sample(factory, tmp_path):
    out = run(factory, tmp_path, SAMPLE.read_bytes(), "sample_assets.xlsx")
    assert out.ok, out.errors
    assert out.asset_count == 22 and out.summary.products_created == 23
    with factory() as s:
        web = s.execute(select(Asset).where(Asset.asset_code == "WEB-001")).scalar_one()
        assert web.zone == Zone.PERIMETER and web.owner_name == "가상담당자01"
        p = web.products[0]
        assert (p.vendor_norm, p.product_norm, p.version_norm) == ("apache", "http_server", "2.4.49")
        # 잘못된 CPE: 오류 기록 후 제품정보 유지 (Level 2 대상)
        mail = s.execute(select(Asset).where(Asset.asset_code == "MAIL-001")).scalar_one()
        assert mail.products[0].cpe_normalized is None and mail.products[0].cpe_error
        # CPE 없는 자산
        was2 = s.execute(select(Asset).where(Asset.asset_code == "WAS-002")).scalar_one()
        assert was2.products[0].cpe_normalized is None and was2.products[0].product_norm == "tomcat"
        # 동일 Asset ID 다중 제품
        lnx = s.execute(select(Asset).where(Asset.asset_code == "LNX-001")).scalar_one()
        assert len(lnx.products) == 2
        # 업데이트 필드 CPE
        vc = s.execute(select(Asset).where(Asset.asset_code == "VC-001")).scalar_one()
        assert vc.products[0].update_norm == "update1d"
        # 수집 대상 산출(Phase 2 연계)
        keys = {t.product_key for t in build_targets(s)}
        assert {"a:apache:http_server", "o:fortinet:fortios", "a:apache:tomcat"} <= keys
        assert s.execute(select(AuditLog).where(AuditLog.action == "EXCEL_UPLOAD",
                                                AuditLog.result == "success")).scalar_one()
        assert s.execute(select(ImportBatch)).scalar_one().original_filename == "sample_assets.xlsx"


@pytest.mark.parametrize("col,value,msg", [
    (3, "999.1.1.1", "IP"), (2, "DMZ", "경계면"), (8, "최상", "중요도"),
    (0, "SRV 01;DROP", "Asset ID"), (1, "", "필수"),
])
def test_invalid_rows_rejected_without_db_change(factory, tmp_path, col, value, msg):
    row = list(ROW)
    row[col] = value
    out = run(factory, tmp_path, xlsx([row]))
    assert not out.ok and out.errors
    assert any(msg in (e.message + e.column) for e in out.errors)
    assert count(factory, Asset) == 0
    with factory() as s:
        assert s.execute(select(AuditLog).where(AuditLog.result == "failure")).scalar_one()


def test_one_bad_row_rejects_whole_file(factory, tmp_path):
    bad = list(ROW)
    bad[0], bad[3] = "SRV-02", "not-an-ip"
    out = run(factory, tmp_path, xlsx([ROW, bad]))
    assert not out.ok and count(factory, Asset) == 0


def test_missing_required_column(factory, tmp_path):
    out = run(factory, tmp_path, xlsx([ROW[:10]], header=HEAD[:10]))
    assert not out.ok and "부서" in out.message


def test_formula_cell_rejected(factory, tmp_path):
    out = run(factory, tmp_path, xlsx([ROW], formula_cells=[("J2", '=HYPERLINK("http://evil","x")')]))
    assert not out.ok and any("수식" in e.message for e in out.errors)


def test_formula_like_text_is_stored_as_plain_data(factory, tmp_path):
    wb = Workbook()
    ws = wb.active
    ws.append(HEAD)
    ws.append(ROW)
    ws["J2"].value = "=1+1"
    ws["J2"].data_type = "s"     # 텍스트로 저장된 '=' 시작 값
    buf = io.BytesIO()
    wb.save(buf)
    out = run(factory, tmp_path, buf.getvalue())
    assert out.ok
    with factory() as s:
        assert s.execute(select(Asset)).scalar_one().owner_name == "=1+1"   # 실행 아닌 데이터 (Export 시 무해화)


def test_inconsistent_duplicate_asset_rows(factory, tmp_path):
    r2 = list(ROW)
    r2[5], r2[7], r2[3] = "Log4j", "cpe:2.3:a:apache:log4j:2.14.1:*:*:*:*:*:*:*", "192.0.2.99"
    out = run(factory, tmp_path, xlsx([ROW, r2]))
    assert not out.ok and any("다른 행" in e.message for e in out.errors)


def test_malicious_upload_rejected_and_audited(factory, tmp_path):
    out = run(factory, tmp_path, b"MZ\x90\x00evil", "../../payload.xlsx")
    assert not out.ok
    with factory() as s:
        log = s.execute(select(AuditLog).where(AuditLog.action == "EXCEL_UPLOAD")).scalar_one()
        assert log.result == "failure" and log.after["file"] == "payload.xlsx"
    up = tmp_path / "uploads"
    assert not up.exists() or not any(up.iterdir())   # 검증 실패 파일은 저장하지 않음


def test_reimport_updates_versions_deactivates_and_flags_revalidation(factory, tmp_path):
    other = list(ROW)
    other[0], other[3] = "SRV-02", "192.0.2.2"
    assert run(factory, tmp_path, xlsx([ROW, other])).ok
    # 매핑이 존재하는 상태를 만든다
    with factory() as s:
        p = s.execute(select(AssetProduct).join(Asset).where(Asset.asset_code == "SRV-01")).scalar_one()
        s.add(Vulnerability(cve_id="CVE-2020-1938", cvss_score=9.8))
        s.flush()
        s.add(AssetVulnerability(asset_id=p.asset_id, asset_product_id=p.id, cve_id="CVE-2020-1938",
                                 assessment_state=AssessmentState.ASSESSED, match_type=MatchType.CPE_RANGE,
                                 owner="가상담당"))
        s.commit()
    # 두 번째 업로드: SRV-01 버전 변경 + 담당자 변경, SRV-02 제외
    upd = list(ROW)
    upd[6], upd[7], upd[9] = "9.0.31", "cpe:2.3:a:apache:tomcat:9.0.31:*:*:*:*:*:*:*", "새담당"
    out = run(factory, tmp_path, xlsx([upd]))
    assert out.ok
    s_ = out.summary
    assert (s_.assets_updated, s_.assets_deactivated, s_.products_updated, s_.revalidation_flagged) == (1, 1, 1, 1)
    with factory() as s:
        av = s.execute(select(AssetVulnerability)).scalar_one()
        assert av.needs_revalidation and av.owner == "새담당"
        srv2 = s.execute(select(Asset).where(Asset.asset_code == "SRV-02")).scalar_one()
        assert srv2.is_active is False                       # 삭제하지 않고 비활성화
        assert s.execute(select(AuditLog).where(AuditLog.action == "OWNER_CHANGE")).scalar_one()
        keys = {t.product_key for t in build_targets(s)}
        assert keys == {"a:apache:tomcat"}
    # 동일 파일 재업로드: 변경 없음
    out2 = run(factory, tmp_path, xlsx([upd]))
    assert (out2.summary.assets_updated, out2.summary.products_updated) == (0, 0)


def test_normalize_token():
    assert normalize_token(" Apache HTTP Server ") == "apache_http_server"
    assert normalize_token("Red Hat / Enterprise") == "red_hat_enterprise"
    assert normalize_token("<script>") == "script"
    assert normalize_token("") is None
