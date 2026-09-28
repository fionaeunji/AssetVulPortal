"""수집 → 자산 → 매핑 → 판정 → 기한 → 승인 통합 테스트 (sample_assets.xlsx + 실제 NVD 레코드 Fixture)."""
from __future__ import annotations

from datetime import date, timedelta
from pathlib import Path

import pytest
from sqlalchemy import select
from sqlalchemy.orm import sessionmaker

from app.config.settings import Settings
from app.models import (
    Asset,
    AssetProduct,
    AssetProductMapping,
    AssetVulnerability,
    AuditLog,
    MappingCandidate,
    PolicyVersion,
    VulnerabilityAssessment,
)
from app.models.enums import AssessmentState, CandidateStatus, CollectionTrigger, MatchType
from app.services.asset_importer import import_asset_upload
from app.services.mapping_service import decide_candidate, run_mapping, run_mapping_with_active_policy
from app.services.nvd_client import NvdClient
from app.services.policy_loader import apply_policy_file, get_active_policy
from app.services.vulnerability_collector import run_online_collection
from tests.fakes import FakeExternal

ROOT = Path(__file__).resolve().parents[2]
SAMPLE = ROOT / "sample_data" / "sample_assets.xlsx"
EPSS = {"CVE-2021-41773": (0.94, 0.99), "CVE-2021-42013": (0.94, 0.99), "CVE-2020-1938": (0.20, 0.97),
        "CVE-2024-38063": (0.05, 0.9), "CVE-2023-4863": (0.9, 0.99), "CVE-2024-21762": (0.4, 0.99),
        "CVE-2022-21449": (0.3, 0.9), "CVE-2023-39417": (0.01, 0.5)}


@pytest.fixture()
def env(engine, tmp_path):
    factory = sessionmaker(bind=engine, expire_on_commit=False)
    settings = Settings(_env_file=None, collector_mode="online", data_dir=tmp_path / "data")
    out = import_asset_upload(factory, filename="sample_assets.xlsx", data=SAMPLE.read_bytes(), actor="op1",
                              actor_role="operator", upload_dir=settings.upload_dir, max_bytes=5 << 20,
                              policy_file=settings.policy_file)
    assert out.ok
    run_online_collection(CollectionTrigger.MANUAL, "op1", factory=factory, settings=settings,
                          transport=FakeExternal(epss=EPSS).transport(),
                          nvd_client_factory=lambda http: NvdClient(http, sleep=lambda s: None))
    return factory, settings


def _av(s, asset_code, cve):
    return s.execute(select(AssetVulnerability).join(Asset, Asset.id == AssetVulnerability.asset_id)
                     .where(Asset.asset_code == asset_code, AssetVulnerability.cve_id == cve)).scalar_one_or_none()


@pytest.mark.parametrize("asset,cve,mtype", [
    ("WEB-001", "CVE-2021-41773", MatchType.CPE_EXACT),
    ("WAS-001", "CVE-2020-1938", MatchType.CPE_RANGE),
    ("FW-001", "CVE-2024-21762", MatchType.CPE_RANGE),
    ("FW-001", "CVE-2023-27997", MatchType.CPE_RANGE),          # versionEndIncluding 7.2.4 경계
    ("RT-001", "CVE-2023-20198", MatchType.CPE_RANGE),
    ("VC-001", "CVE-2021-21972", MatchType.CPE_EXACT),
    ("WIN-001", "CVE-2024-38063", MatchType.CPE_RANGE),
    ("DB-001", "CVE-2023-39417", MatchType.CPE_RANGE),
    ("APP-002", "CVE-2022-21449", MatchType.CPE_EXACT),
    ("APP-003", "CVE-2021-44228", MatchType.CPE_RANGE),
    ("APP-004", "CVE-2024-3094", MatchType.CPE_EXACT),
    ("PC-001", "CVE-2023-4863", MatchType.CPE_RANGE),
])
def test_expected_vulnerable_mappings(env, asset, cve, mtype):
    factory, _ = env
    with factory() as s:
        av = _av(s, asset, cve)
        assert av is not None and av.match_type == mtype and not av.needs_revalidation
        assert av.match_evidence["criteria"] and av.match_evidence["rule"]
        assert av.owner and av.match_confidence is None


@pytest.mark.parametrize("asset,cve", [
    ("WEB-002", "CVE-2021-41773"), ("FW-002", "CVE-2024-21762"), ("WIN-002", "CVE-2024-38063"),
    ("PC-002", "CVE-2023-4863"),
])
def test_patched_versions_not_mapped(env, asset, cve):
    factory, _ = env
    with factory() as s:
        assert _av(s, asset, cve) is None


def test_no_cpe_assets_never_auto_vulnerable(env):
    factory, _ = env
    with factory() as s:
        for code in ("WAS-002", "APP-001", "ESX-001", "LNX-002", "DB-002", "MAIL-001"):
            a = s.execute(select(Asset).where(Asset.asset_code == code)).scalar_one()
            assert not s.execute(select(AssetVulnerability).where(AssetVulnerability.asset_id == a.id)).first()


def test_level2_candidates_created_with_rule_scores(env):
    factory, _ = env
    with factory() as s:
        cands = s.execute(select(MappingCandidate).where(MappingCandidate.level == 2)).scalars().all()
        by_asset = {s.get(Asset, s.get(AssetProduct, c.asset_product_id).asset_id).asset_code: c for c in cands}
        assert {"WAS-002", "APP-001"} <= set(by_asset)
        was2 = by_asset["WAS-002"]
        assert was2.proposed_cpe == "cpe:2.3:a:apache:tomcat:8.5.50:*:*:*:*:*:*:*"
        assert was2.confidence == sum(was2.breakdown["rules"].values()) == 100   # 수집 후 점수 갱신
        assert by_asset["APP-001"].proposed_cpe.startswith("cpe:2.3:a:oracle:jdk:1.8.0:update_401")
        assert "MAIL-001" not in by_asset


def test_level3_review_when_update_unknown(env):
    factory, _ = env
    with factory() as s:
        c = s.execute(select(MappingCandidate).where(MappingCandidate.level == 3,
                                                     MappingCandidate.cve_id == "CVE-2021-21972")).scalar_one()
        p = s.get(AssetProduct, c.asset_product_id)
        assert s.get(Asset, p.asset_id).asset_code == "VC-002"
        assert c.confidence is None and "update" in c.reason
        assert _av(s, "VC-002", "CVE-2021-21972") is None             # 자동으로 취약 확정하지 않음


def test_severity_and_deadlines(env):
    factory, _ = env
    with factory() as s:
        web = _av(s, "WEB-001", "CVE-2021-41773")          # 9.8 / 0.94 / 경계면
        assert web.severity == "긴급" and web.due_at - web.detected_at == timedelta(hours=72)
        assert web.initial_due_at == web.due_at
        was = _av(s, "WAS-001", "CVE-2020-1938")           # 9.8 / 0.20 / 내부
        assert was.severity == "우선" and was.due_at - was.detected_at == timedelta(days=45)
        win = _av(s, "WIN-001", "CVE-2024-38063")          # 9.8 / 0.05 / 내부
        assert win.severity == "주의" and win.due_at - win.detected_at == timedelta(days=90)
        log4j = _av(s, "APP-003", "CVE-2021-44228")        # CVSS 10, EPSS 없음
        assert log4j.assessment_state == AssessmentState.EPSS_PENDING and log4j.severity is None
        snap = s.execute(select(VulnerabilityAssessment).where(
            VulnerabilityAssessment.asset_vulnerability_id == web.id)).scalar_one()
        assert snap.cvss_score_used == 9.8 and snap.epss_initial_used == 0.94 and "72" in snap.calc_method


def test_epss_pending_resolved_on_first_epss(env):
    factory, settings = env
    fake = FakeExternal(epss={**EPSS, "CVE-2021-44228": (0.97, 0.99)}, epss_date=date(2026, 9, 28))
    run_online_collection(CollectionTrigger.MANUAL, "op1", factory=factory, settings=settings,
                          transport=fake.transport(),
                          nvd_client_factory=lambda http: NvdClient(http, sleep=lambda s: None))
    with factory() as s:
        av = _av(s, "APP-003", "CVE-2021-44228")
        assert av.assessment_state == AssessmentState.ASSESSED and av.severity == "긴급"
        assert av.due_at - av.detected_at == timedelta(days=45)       # 내부 자산, 탐지일 기준


def test_rerun_is_idempotent(env):
    factory, settings = env
    with factory() as s:
        n_av = len(s.execute(select(AssetVulnerability)).all())
        n_snap = len(s.execute(select(VulnerabilityAssessment)).all())
        n_cand = len(s.execute(select(MappingCandidate)).all())
        sm = run_mapping_with_active_policy(s, actor="t", policy_file=settings.policy_file)
        s.commit()
        assert sm.vulnerable_new == 0 and sm.assessments_written == 0
        assert len(s.execute(select(AssetVulnerability)).all()) == n_av
        assert len(s.execute(select(VulnerabilityAssessment)).all()) == n_snap
        assert len(s.execute(select(MappingCandidate)).all()) == n_cand


def test_approve_level2_mapping_is_reused_and_audited(env):
    factory, settings = env
    with factory() as s:
        c = s.execute(select(MappingCandidate).where(
            MappingCandidate.level == 2,
            MappingCandidate.proposed_cpe == "cpe:2.3:a:apache:tomcat:8.5.50:*:*:*:*:*:*:*")).scalar_one()
        decide_candidate(s, c.id, approve=True, actor="op1", actor_role="operator", reason="자산 담당 확인")
        sm = run_mapping_with_active_policy(s, actor="op1", policy_file=settings.policy_file)
        s.commit()
        av = _av(s, "WAS-002", "CVE-2020-1938")          # 8.5.0 ≤ 8.5.50 < 8.5.51
        assert av and av.match_type == MatchType.APPROVED_MAPPING and av.mapping_approved_by == "op1"
        assert av.match_evidence["mapping_approved_by"] == "op1"
        m = s.execute(select(AssetProductMapping)).scalar_one()
        assert (m.vendor_norm, m.product_norm, m.cpe_vendor, m.cpe_product) == ("apache", "tomcat", "apache", "tomcat")
        assert s.execute(select(AuditLog).where(AuditLog.action == "MAPPING_APPROVE")).scalar_one()
    # 동일 제품의 신규 자산 → 승인된 매핑 자동 재사용
    import io

    from openpyxl import load_workbook
    wb = load_workbook(SAMPLE)
    ws = wb.active
    ws.append(["WAS-003", "was-app-03", "내부", "192.0.2.23", "Apache", "Tomcat", "8.5.40", "", "중",
               "가상담당자02", "가상개발팀", ""])
    buf = io.BytesIO()
    wb.save(buf)
    out = import_asset_upload(factory, filename="v2.xlsx", data=buf.getvalue(), actor="op1",
                              actor_role="operator", upload_dir=settings.upload_dir, max_bytes=5 << 20,
                              policy_file=settings.policy_file)
    assert out.ok
    with factory() as s:
        av = _av(s, "WAS-003", "CVE-2020-1938")
        assert av and av.match_type == MatchType.APPROVED_MAPPING


def test_reject_level2_not_reproposed(env):
    factory, settings = env
    with factory() as s:
        c = s.execute(select(MappingCandidate).where(
            MappingCandidate.level == 2, MappingCandidate.proposed_cpe.like("cpe:2.3:a:oracle:jdk:%"))).scalar_one()
        decide_candidate(s, c.id, approve=False, actor="op1", actor_role="operator")
        s.commit()
        assert s.execute(select(AuditLog).where(AuditLog.action == "MAPPING_REJECT")).scalar_one()
        run_mapping_with_active_policy(s, actor="t", policy_file=settings.policy_file)
        pend = s.execute(select(MappingCandidate).where(
            MappingCandidate.status == CandidateStatus.PENDING,
            MappingCandidate.proposed_cpe.like("cpe:2.3:a:oracle:jdk:%"))).all()
        assert pend == []


def test_approve_level3_confirms_with_manual_evidence(env):
    factory, settings = env
    with factory() as s:
        c = s.execute(select(MappingCandidate).where(MappingCandidate.level == 3,
                                                     MappingCandidate.cve_id == "CVE-2021-21972")).scalar_one()
        decide_candidate(s, c.id, approve=True, actor="admin", actor_role="admin", reason="6.7 U3 이하 확인")
        run_mapping_with_active_policy(s, actor="t", policy_file=settings.policy_file)
        s.commit()
        av = _av(s, "VC-002", "CVE-2021-21972")
        assert av.match_evidence["manual"] and av.mapping_approved_by == "admin" and not av.needs_revalidation
        with pytest.raises(Exception):
            decide_candidate(s, c.id, approve=True, actor="admin", actor_role="admin")   # 중복 처리 거부


def test_policy_change_preserves_history(env):
    factory, settings = env
    with factory() as s:
        web = _av(s, "WEB-001", "CVE-2021-41773")
        first_due = web.initial_due_at
        text = settings.policy_file.read_text(encoding="utf-8").replace(
            'version: "2026.09-01"', 'version: "2026.10-test"').replace("hours: 72", "hours: 48")
        newfile = settings.data_dir / "policy2.yaml"
        newfile.write_text(text, encoding="utf-8")
        apply_policy_file(s, newfile, actor="admin")
        pv, policy = get_active_policy(s)
        run_mapping(s, pv=pv, policy=policy)
        s.commit()
        s.refresh(web)
        assert web.due_at - web.detected_at == timedelta(hours=48)
        assert web.initial_due_at == first_due                         # 최초 기한 보존
        snaps = s.execute(select(VulnerabilityAssessment).where(
            VulnerabilityAssessment.asset_vulnerability_id == web.id)
            .order_by(VulnerabilityAssessment.id)).scalars().all()
        versions = [s.get(PolicyVersion, x.policy_version_id).version for x in snaps]
        assert versions == ["2026.09-01", "2026.10-test"]              # 과거 판정 보존
        assert s.execute(select(AuditLog).where(AuditLog.action == "POLICY_CHANGE")).scalars().all()


def test_version_upgrade_flags_revalidation(env):
    import io

    from openpyxl import load_workbook
    factory, settings = env
    wb = load_workbook(SAMPLE)
    ws = wb.active
    for row in ws.iter_rows(min_row=2):
        if row[0].value == "WEB-001":
            row[6].value = "2.4.62"
            row[7].value = "cpe:2.3:a:apache:http_server:2.4.62:*:*:*:*:*:*:*"
    buf = io.BytesIO()
    wb.save(buf)
    assert import_asset_upload(factory, filename="v3.xlsx", data=buf.getvalue(), actor="op1",
                               actor_role="operator", upload_dir=settings.upload_dir, max_bytes=5 << 20,
                               policy_file=settings.policy_file).ok
    with factory() as s:
        av = _av(s, "WEB-001", "CVE-2021-41773")
        assert av is not None and av.needs_revalidation     # 삭제하지 않고 재검증 필요 표시
