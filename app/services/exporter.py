"""취약점 현황 Excel(.xlsx) Export — Formula(CSV/Excel) Injection 방어.

외부 데이터(CVE 설명)·사용자 입력(Excel 자산명, 담당자 등)은 신뢰하지 않는다.
- 문자열이 수식 시작 문자( = + - @ , 탭, CR, LF, 전각 ＝＋－＠ )로 시작하면 앞에 작은따옴표(')를 붙여 텍스트로 고정
- 모든 문자열 셀의 data_type 을 's'(문자열)로 강제 → openpyxl 이 '=' 시작 값을 수식으로 저장하지 않음
- 숫자·날짜는 실제 숫자/날짜 타입으로 기록 (문자열 변환 없음)
- 누락값은 빈 셀로 기록 ('-' 같은 표기를 쓰지 않아 불필요한 무해화 방지)
"""
from __future__ import annotations

import io
from datetime import datetime
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

KST = ZoneInfo("Asia/Seoul")
FORMULA_PREFIXES = ("=", "+", "-", "@", "\t", "\r", "\n", "＝", "＋", "－", "＠")
MAX_CELL_CHARS = 32000     # Excel 셀 최대 32,767자
MAX_EXPORT_ROWS = 50_000

COLUMNS = [
    ("등급", 8), ("CVE", 16), ("자산 ID", 11), ("자산명", 16), ("IP", 15), ("제품", 22), ("Version", 16),
    ("CVSS", 6), ("CVSS 버전", 8), ("CVSS Vector", 44), ("Initial EPSS", 11), ("Current EPSS", 11),
    ("KEV", 5), ("자산구분", 8), ("중요도", 6), ("담당자", 12), ("부서", 12), ("탐지일", 17), ("조치기한", 17),
    ("최초 조치기한", 17), ("남은시간(시간)", 12), ("기한초과", 8), ("상태", 9), ("매핑방식", 17),
    ("매핑근거", 50), ("재검증 필요", 9), ("CVE 설명", 60),
]


def sanitize_text(value: str) -> str:
    """Formula Injection 무해화."""
    v = value[:MAX_CELL_CHARS]
    if v.startswith(FORMULA_PREFIXES):
        return "'" + v
    return v


def _put(ws, row: int, col: int, value):
    cell = ws.cell(row=row, column=col)
    if value is None:
        return cell
    if isinstance(value, bool):
        cell.value = "Y" if value else "N"
        cell.data_type = "s"
    elif isinstance(value, (int, float)):
        cell.value = value
    elif isinstance(value, datetime):
        cell.value = value.astimezone(KST).replace(tzinfo=None)
        cell.number_format = "yyyy-mm-dd hh:mm"
    else:
        cell.value = sanitize_text(str(value))
        cell.data_type = "s"          # 수식으로 해석되지 않도록 문자열 타입 강제
    return cell


def _evidence_text(ev: dict | None) -> str | None:
    if not ev:
        return None
    keys = ("criteria", "rule", "match_criteria_range", "platform_conditions", "asset_cpe",
            "mapping_approved_by", "review_reason", "approved_by")
    parts = []
    for k in keys:
        v = ev.get(k)
        if v:
            parts.append(f"{k}: {'; '.join(map(str, v)) if isinstance(v, list) else v}")
    return " | ".join(parts) or None


def build_workbook(rows, *, now: datetime, exported_by: str, policy_version: str,
                   filter_desc: str, kpi: dict, closed_statuses: set[str]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "취약점 현황"
    hdr_fill = PatternFill("solid", fgColor="DDEBF7")
    for c, (name, width) in enumerate(COLUMNS, start=1):
        cell = _put(ws, 1, c, name)
        cell.font = Font(bold=True)
        cell.fill = hdr_fill
        ws.column_dimensions[get_column_letter(c)].width = width
    r = 1
    for av, a, p, v in rows[:MAX_EXPORT_ROWS]:
        r += 1
        closed = av.status.value in closed_statuses
        remain_h = round((av.due_at - now).total_seconds() / 3600, 1) if av.due_at and not closed else None
        values = [
            av.severity or ("EPSS 대기" if av.assessment_state.value == "epss_pending" else None),
            v.cve_id, a.asset_code, a.name, a.ip, p.product_raw or p.product_norm, p.version_norm,
            v.cvss_score, v.cvss_version, v.cvss_vector, v.epss_initial, v.epss_current,
            bool(v.kev), a.zone.value, a.criticality.value, av.owner, a.department,
            av.detected_at, av.due_at, av.initial_due_at, remain_h,
            bool(remain_h is not None and remain_h < 0), av.status.value, av.match_type.value,
            _evidence_text(av.match_evidence), bool(av.needs_revalidation), v.description,
        ]
        for c, val in enumerate(values, start=1):
            _put(ws, r, c, val)
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNS))}{max(r, 1)}"
    for row in ws.iter_rows(min_row=2, min_col=25, max_col=27):
        for cell in row:
            cell.alignment = Alignment(wrap_text=False, vertical="top")

    info = wb.create_sheet("요약")
    items = [("보고서", "IT자산 취약점 현황"), ("생성일시(KST)", now), ("생성자", exported_by),
             ("정책 버전", policy_version), ("필터", filter_desc or "기본(미종결·관리대상)"),
             ("행 수", min(len(rows), MAX_EXPORT_ROWS)), ("전체 자산", kpi.get("total_assets")),
             ("취약 자산", kpi.get("vulnerable_assets"))]
    items += [(f"{k} 취약점", n) for k, n in (kpi.get("by_severity") or {}).items()]
    items += [("EPSS 대기", kpi.get("epss_pending")), ("조치기한 초과", kpi.get("overdue")),
              ("비고", "시각은 Asia/Seoul 기준. SOURCE: NVD/FIRST EPSS/CISA KEV, 등급·기한은 시스템 계산값")]
    for i, (k, val) in enumerate(items, start=1):
        _put(info, i, 1, k).font = Font(bold=True)
        _put(info, i, 2, val)
    info.column_dimensions["A"].width = 18
    info.column_dimensions["B"].width = 60
    if len(rows) > MAX_EXPORT_ROWS:
        _put(info, len(items) + 1, 1, "경고")
        _put(info, len(items) + 1, 2, f"최대 {MAX_EXPORT_ROWS}행까지만 포함됨 (전체 {len(rows)}행)")
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def export_filename(now: datetime) -> str:
    """서버 생성 파일명 (사용자 입력 미사용)."""
    return f"vuln_report_{now.astimezone(KST).strftime('%Y%m%d_%H%M')}.xlsx"
